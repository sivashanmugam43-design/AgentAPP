"""Pipeline, store and HTTP client tests. No network: fake fetcher, mock HTTP transport."""

import asyncio
import csv
import json

import httpx
import pytest

from config.settings import settings
from src.temples import pipeline
from src.temples.scraper import FetchError, PoliteClient
from src.temples.store import RECORD_COLUMNS, Store


def page(temples, next_page=None):
    links = "".join(f'<a class="newsdetails" href="new.php?id={tid}">{name}</a>' for tid, name in temples)
    nxt = f'<a href="district_temple.php?id=46&Page={next_page}">Next &gt;&gt;</a>' if next_page else ""
    return f"<html><body>{links}{nxt}</body></html>"


def temple_html(name, moolavar, english=False, extra_rows=""):
    en = '<a class="click" href="en/new_en.php?id=1">English Version</a>' if english else ""
    return f"""<html><body><div class="TabbedPanelsContent">{en}
      <span class="topic"><span id="tknameLabel">{name}</span></span>
      <table><tr><td class="style8">மூலவர்</td><td>:</td><td class="style5">{moolavar}</td></tr>
             <tr><td class="style8">ஊர்</td><td>:</td><td class="style5">{"மயிலாப்பூர்" if moolavar else ""}</td></tr>{extra_rows}</table>
      <table><tr><td><img src="//x/images/l.gif"></td><td><span class="subhead">திறக்கும் நேரம்:</span></td></tr></table>
      <table><tr><td class="newsdetails">காலை 6 மணி</td></tr></table>
    </div></body></html>"""


def english_html():
    return """<html><body><div class="TabbedPanelsContent">
      <span class="topic">Sri Test temple</span>
      <table><tr><td class="style8">Moolavar</td><td>:</td><td class="style5">Kapaleeswarar</td></tr>
             <tr><td class="style8">City</td><td>:</td><td class="style5">Mylapore</td></tr></table>
    </div></body></html>"""


class FakeFetcher:
    def __init__(self, pages, fail=()):
        self.pages, self.fail, self.calls = pages, set(fail), []

    async def get_text(self, url):
        self.calls.append(url)
        if url in self.fail:
            raise FetchError("HTTP 503 (after 5 attempts)")
        return self.pages[url]


@pytest.fixture
def store(tmp_path):
    return Store(tmp_path / "out", tmp_path / "progress")


DISTRICT = {"district_id": "46", "district_name": "சென்னை"}


def test_stage2_follows_next_and_resumes_after_failure(store):
    p1, p2, p3 = (settings.district_url("46", n) for n in (1, 2, 3))
    pages = {p1: page([("1", "A"), ("2", "B")], next_page=2),
             p2: page([("2", "B"), ("3", "C")], next_page=3),   # temple 2 repeated: deduped
             p3: page([("4", "D")])}
    fetcher = FakeFetcher(pages, fail={p2})
    asyncio.run(pipeline.stage2_index(fetcher, store, [DISTRICT]))
    assert [r["temple_id"] for r in store.load_index()] == ["1", "2"]
    assert store.load_stage2_state()["46"]["last_page"] == 1
    assert "\t" + p2 + "\tHTTP 503" in store.failed_path.read_text(encoding="utf-8")

    fetcher = FakeFetcher(pages)  # second run: page 2 works again; page 1 is not re-fetched
    asyncio.run(pipeline.stage2_index(fetcher, store, [DISTRICT]))
    assert fetcher.calls == [p2, p3]
    assert [r["temple_id"] for r in store.load_index()] == ["1", "2", "3", "4"]
    assert store.load_stage2_state()["46"] == {"last_page": 3, "complete": True, "temples": 4}

    fetcher = FakeFetcher(pages)  # third run: district complete, nothing fetched
    asyncio.run(pipeline.stage2_index(fetcher, store, [DISTRICT]))
    assert fetcher.calls == []


def test_index_csv_has_bom_once_and_tamil_intact(store):
    store.append_index([{"district_id": "46", "district_name": "சென்னை", "temple_id": "1",
                         "temple_name": "அருள்மிகு", "url": "u1"}])
    store.append_index([{"district_id": "46", "district_name": "சென்னை", "temple_id": "2",
                         "temple_name": "திருக்கோயில்", "url": "u2"}])
    raw = store.index_path.read_bytes()
    assert raw.startswith(b"\xef\xbb\xbf") and raw.count(b"\xef\xbb\xbf") == 1
    assert [r["temple_name"] for r in store.load_index()] == ["அருள்மிகு", "திருக்கோயில்"]


def test_stage3_site_english_and_resume(store):
    store.append_index([
        {"district_id": "46", "district_name": "சென்னை", "temple_id": tid, "temple_name": "x",
         "url": settings.temple_url(tid)} for tid in ("1", "2", "3")])
    pages = {
        settings.temple_url("1"): temple_html("அருள்மிகு கபாலீஸ்வரர் திருக்கோயில்", "கபாலீசுவரர்", english=True),
        settings.temple_url_en("1"): english_html(),
        settings.temple_url("2"): temple_html("அருள்மிகு திருக்கோயில்", ""),  # empty template
        settings.temple_url("3"): temple_html("அருள்மிகு ராமர் திருக்கோயில்", "ராமர்"),
    }
    stats = asyncio.run(pipeline.stage3_temples(FakeFetcher(pages), store))
    assert (stats["ok"], stats["failed"], stats["english_page"]) == (2, 1, 1)

    recs = {r["temple_id"]: r for r in store.load_records()}
    assert set(recs) == {"1", "3"}
    r1 = recs["1"]
    assert (r1["moolavar_en"], r1["moolavar_source_en"]) == ("Kapaleeswarar", "site")
    # On the Tamil page but blank in English: left for translate.py
    assert (r1["timings_ta"], r1["timings_en"], r1["timings_source_en"]) == ("காலை 6 மணி", "", "")
    assert (r1["urchavar_ta"], r1["urchavar_en"], r1["urchavar_source_en"]) == ("", "", "")
    r3 = recs["3"]
    assert r3["english_page"] == "no" and (r3["name_en"], r3["name_source_en"]) == ("", "")
    assert "id=2\tTempleNotFound" in store.failed_path.read_text(encoding="utf-8")

    fetcher = FakeFetcher(pages)  # rerun: only the failed temple is tried again
    asyncio.run(pipeline.stage3_temples(fetcher, store))
    assert fetcher.calls == [settings.temple_url("2")]


def test_stage3_district_filter_and_limit(store):
    store.append_index([{"district_id": d, "district_name": "", "temple_id": t, "temple_name": "", "url": ""}
                        for d, t in (("46", "1"), ("46", "2"), ("46", "3"), ("47", "4"))])
    pages = {settings.temple_url(t): temple_html("n", "m") for t in "1234"}
    fetcher = FakeFetcher(pages)
    asyncio.run(pipeline.stage3_temples(fetcher, store, district_id="46", limit=2))
    assert sorted(fetcher.calls) == [settings.temple_url("1"), settings.temple_url("2")]


def test_append_record_recovers_from_cut_line(store):
    store.records_path.write_text('{"temple_id": "1"}\n{"temple_id": "2", "na', encoding="utf-8")
    store.append_record({"temple_id": "3", "name_ta": "கோயில்"})
    assert [r["temple_id"] for r in store.load_records()] == ["1", "3"]


def test_export_csv_and_xlsx(store):
    from openpyxl import load_workbook
    rec = {c: "" for c in RECORD_COLUMNS}
    rec.update(temple_id="628", district_id="46", name_ta="அருள்மிகு கபாலீஸ்வரர் திருக்கோயில்",
               name_en="Sri Kapaleeswarar temple", name_source_en="site",
               gallery=[{"url": "u", "caption_ta": "அம்மன்"}], history_ta="அ" * 40000)
    paths = store.export([rec])
    raw = paths["csv"].read_bytes()
    assert raw.startswith(b"\xef\xbb\xbf")
    rows = list(csv.DictReader(paths["csv"].open(encoding="utf-8-sig", newline="")))
    assert rows[0]["name_ta"] == "அருள்மிகு கபாலீஸ்வரர் திருக்கோயில்"
    assert json.loads(rows[0]["gallery"])[0]["caption_ta"] == "அம்மன்"
    assert len(rows[0]["history_ta"]) == 40000
    ws = load_workbook(paths["xlsx"]).active
    header = [c.value for c in ws[1]]
    assert header == RECORD_COLUMNS
    assert ws.cell(2, header.index("name_ta") + 1).value == "அருள்மிகு கபாலீஸ்வரர் திருக்கோயில்"
    assert len(ws.cell(2, header.index("history_ta") + 1).value) == 32767


# --- HTTP client ---------------------------------------------------------------------------------

def run_client(handler, **kw):
    async def go():
        async with PoliteClient(delay_min=0, delay_max=0, backoff_base=0.01,
                                transport=httpx.MockTransport(handler), **kw) as client:
            return await client.get_text("https://temple.dinamalar.com/new.php?id=1")
    return asyncio.run(go())


def test_client_retries_5xx_and_429_then_decodes_utf8():
    responses = iter([httpx.Response(503), httpx.Response(429, headers={"Retry-After": "0"}),
                      httpx.Response(200, content="அருள்மிகு".encode("utf-8"),
                                     headers={"Content-Type": "text/html; charset=ISO-8859-1"})])
    assert run_client(lambda request: next(responses)) == "அருள்மிகு"


def test_client_gives_up_after_max_attempts_and_fails_fast_on_404():
    calls = []
    def flaky(request):
        calls.append(1)
        raise httpx.ConnectTimeout("timed out")
    with pytest.raises(FetchError, match="ConnectTimeout.*after 3 attempts"):
        run_client(flaky, max_attempts=3)
    assert len(calls) == 3
    calls.clear()
    def missing(request):
        calls.append(1)
        return httpx.Response(404)
    with pytest.raises(FetchError, match="HTTP 404"):
        run_client(missing)
    assert len(calls) == 1
